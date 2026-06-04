import os
from datetime import datetime
from dataclasses import dataclass
from enum import Enum
from typing import List
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.drawing.image import Image as XLImage
from config.settings import REPORT_DIR


class CashFlowResult(Enum):
    PASS = "PASS"
    FAIL = "FAIL"


@dataclass
class SpinRecord:
    index: int
    timestamp: str
    balance_before: float
    bet: float
    win: float
    balance_after: float
    result: CashFlowResult
    screenshot_path: str = ""
    failure_reason: str = ""


def verify_cashflow(record: SpinRecord) -> CashFlowResult:
    expected = round(record.balance_before - record.bet + record.win, 2)
    actual = round(record.balance_after, 2)
    return CashFlowResult.PASS if abs(expected - actual) < 0.5 else CashFlowResult.FAIL


def generate_excel_report(records: List[SpinRecord]) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Slot Test Report"

    headers = ["項次", "測試時間", "玩家資產(前)", "押注", "遊戲贏分", "玩家資產(後)", "測試結果", "截圖"]
    header_fill = PatternFill(start_color="4B4B8F", end_color="4B4B8F", fill_type="solid")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    pass_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    fail_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    for row_idx, r in enumerate(records, 2):
        ws.row_dimensions[row_idx].height = 80
        fill = pass_fill if r.result == CashFlowResult.PASS else fail_fill
        values = [r.index, r.timestamp, r.balance_before, r.bet,
                  r.win, r.balance_after, r.result.value]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        if r.screenshot_path and os.path.exists(r.screenshot_path):
            try:
                img = XLImage(r.screenshot_path)
                img.width, img.height = 120, 70
                ws.add_image(img, f"H{row_idx}")
            except Exception:
                ws.cell(row=row_idx, column=8, value=r.screenshot_path)
        elif r.failure_reason:
            cell = ws.cell(row=row_idx, column=8, value=r.failure_reason)
            cell.fill = fill

    col_widths = {'A': 8, 'B': 20, 'C': 15, 'D': 10, 'E': 12, 'F': 15, 'G': 12, 'H': 20}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    filename = f"slot_test_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(REPORT_DIR, filename)
    wb.save(filepath)
    return filepath
